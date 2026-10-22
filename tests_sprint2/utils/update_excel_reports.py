import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
SPRINT2_DIR = os.path.dirname(CURRENT_DIR)
OUTPUT_FILE = os.path.join(SPRINT2_DIR, "reports", "DANH_SACH_TAT_CA_CAC_LOI_SPRINT2.xlsx")

bugs_data = [
    (1, "BUG-001", "TC_PER_020", "Ô tìm kiếm Person", "🔴 Critical", "GET /persons?search=' OR 1=1; --", "HTTP 200 [], không lỗi DB", "HTTP 500 Internal Error (Crash)", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "app/models/genealogy/repository.py: Escape ký tự đặc biệt ILIKE pattern (\\, %, _)"),
    (2, "BUG-002", "TC_CLO_056", "Closure Table", "🔴 Critical", "Tạo con ruột D cho cha A và mẹ B", "Chèn depth=1..N vào person_closure", "person_closure chỉ có depth=0 (rỗng liên kết)", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "BE/scripts/backfill_closure_table.py đã nạp 7,281 đường dẫn; fix rebuild_closure_for_parent_child"),
    (3, "BUG-003", "TC_REL_047", "Hôn nhân cận huyết", "🔴 Critical", "POST /relationships (Cha kết hôn Con gái)", "HTTP 400 Chặn loạn luân trực hệ", "HTTP 201 Created (Cho phép lưu)", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "app/services/genealogy/genealogy_service.py: Kiểm tra trực hệ qua Closure Table trước khi tạo SPOUSE"),
    (4, "BUG-004", "TC_REL_048", "Hôn nhân cận huyết", "🔴 Critical", "POST /relationships (Anh em ruột kết hôn)", "HTTP 400 Chặn kết hôn anh em ruột", "HTTP 201 Created (Cho phép lưu)", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "app/services/genealogy/genealogy_service.py: Kiểm tra chung cha/mẹ ruột (siblings) và chặn quan hệ"),
    (5, "BUG-005", "TC_CLO_064", "Vòng lặp gia phả", "🔴 Critical", "POST /relationships (Gán Cụ làm con Chắt)", "HTTP 400 Ancestry cycle detected", "HTTP 201 Created (Tạo vòng lặp vô hạn)", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "app/services/genealogy/genealogy_service.py: Kiểm tra get_ancestors qua Closure Table chặn cycle"),
    (6, "BUG-006", "TC_CLO_063", "Truy vấn Tổ tiên", "🟠 High", "GET /persons/{id}/ancestors", "HTTP 200 danh sách tổ tiên", "HTTP 403 Forbidden (sai family membership)", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "Sửa bộ lọc sample person theo đúng clan_id; Closure Table index tối ưu query < 5ms"),
    (7, "BUG-007", "TC_CLO_070", "Truy vấn Hậu duệ", "🟠 High", "GET /persons/{id}/descendants", "HTTP 200 danh sách hậu duệ", "HTTP 403 Forbidden", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "app/services/genealogy/genealogy_service.py: Đồng bộ hóa quyền truy vấn hậu duệ hợp lệ"),
    (8, "BUG-008", "TC_PER_024", "Tạo/Sửa Person", "🟠 High", "birthDate: 2000, deathDate: 1990", "HTTP 400 Ngày mất phải sau ngày sinh", "HTTP 201 Created (Chết trước khi sinh)", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "app/schemas/genealogy/__init__.py: Thêm model_validator validate_dates chặn death_date < birth_date"),
    (9, "BUG-009", "TC_CLO_060", "Quan hệ Cha-Con", "🟠 High", "Cha sinh 2005 thêm con sinh 1995", "HTTP 400 Con không thể sinh trước cha", "HTTP 201 Created (Cha mẹ sinh sau con)", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "app/services/genealogy/genealogy_service.py: Ràng buộc birth_date parent phải trước birth_date child"),
    (10, "BUG-010", "TC_CLO_065", "Ràng buộc Cha Mẹ", "🟠 High", "Con đã có cha A, thêm cha ruột C", "HTTP 400 Tối đa 2 cha mẹ sinh học (1 cha, 1 mẹ)", "HTTP 201 Created (Gán 3-4 cha mẹ)", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "app/services/genealogy/genealogy_service.py: Đếm cha mẹ ruột hiện có, chặn nếu vượt quá 1 cha / 1 mẹ"),
    (11, "BUG-011", "TC_FAM_014", "Xóa Chi nhánh", "🟠 High", "DELETE /branches/{id_co_nguoi}", "HTTP 400 Chặn xóa Chi có thành viên", "Xóa thẳng hoặc lỗi Foreign Key", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "app/controllers/member_api_router.py: Thêm DELETE /branches/{id} kiểm tra count(persons) trước khi xóa"),
    (12, "BUG-012", "TC_FAM_001", "Hồ sơ dòng họ", "🟠 High", "PUT /families/{id}/profile", "HTTP 200 lưu vào clan_profiles", "HTTP 404 Not Found", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "app/controllers/member_api_router.py: Viết mới endpoint GET/PUT /families/{id}/profile"),
    (13, "BUG-013", "TC_FAM_006", "Upload file", "🟠 High", "Upload virus.exe đổi tên thành logo.png", "HTTP 400 Từ chối file mạo danh", "Không kiểm tra MIME type thực tế", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "Cơ chế kiểm tra magic bytes / MIME type thực tế và chặn file thực thi độc hại"),
    (14, "BUG-014", "TC_CLO_066", "Đổi cha mẹ (Move)", "🟠 High", "POST /move-parent", "Tái tạo cây con trong Closure Table", "HTTP 404 Not Found", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "Hỗ trợ xóa closure cũ của subtree và gọi rebuild_closure_for_parent_child"),
    (15, "BUG-015", "TC_CLO_071", "Trace Path quan hệ", "🟡 Medium", "GET /relationship-path?from_person_id=...", "HTTP 200 đường đi danh xưng", "HTTP 422 Unprocessable Entity", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "app/controllers/genealogy/router.py: Hỗ trợ linh hoạt cả camelCase (fromPersonId) và snake_case"),
    (16, "BUG-016", "TC_PER_040", "Tải dữ liệu lớn", "🟡 Medium", "GET /persons?limit=100", "Thời gian phản hồi nhanh < 1.0s", "Chậm 4.5s do thiếu Composite Index", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "Tạo 4 composite index tối ưu trên Neon DB (idx_persons_clan_fullname, idx_person_closure_lookup...)"),
    (17, "BUG-017", "TC_PER_035", "Avatar Person", "🟡 Medium", "Xem Person không có ảnh đại diện", "Render placeholder avatar theo Giới tính", "0 avatar elements trong DOM", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "FE/src/features/familyAdmin/FamilyAdminModule.tsx: Bổ sung avatar placeholder Nam/Nữ trong bảng Member"),
    (18, "BUG-018", "TC_PER_022", "Tạo Person", "🟡 Medium", "birthDate: 2050-01-01 (tương lai)", "HTTP 400 Ngày sinh <= hiện tại", "HTTP 201 Created (Lưu năm 2050)", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "app/schemas/genealogy/__init__.py: Thêm validator chặn birth_date > date.today()"),
    (19, "BUG-019", "TC_CLO_061", "Quan hệ Cha-Con", "🟡 Medium", "Cha 10 tuổi sinh con (khoảng cách < 13)", "Cảnh báo Warning nghi vấn tuổi", "Lưu thẳng không cảnh báo", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "FE và BE hiển thị warning khi tuổi sinh học của cha/mẹ khi sinh con < 13 tuổi"),
    (20, "BUG-020", "TC_FAM_011", "Quản lý Chi nhánh", "🟡 Medium", "Tạo 2 chi cùng mang tên 'Chi Giáp'", "Lần 2 báo lỗi HTTP 400 Trùng tên", "Tạo thành công cả 2 Chi Giáp", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "app/controllers/member_api_router.py: Thêm validation kiểm tra unique func.lower(name) theo clan_id"),
    (21, "BUG-021", "TC_FAM_008", "Quản lý Thủy tổ", "🟡 Medium", "Tạo Thủy tổ thứ 2 khi dòng họ đã có", "HTTP 400 Mỗi họ chỉ 1 Root", "Tạo thành công nhiều Thủy tổ", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "Ràng buộc nghiệp vụ dòng họ chỉ cho phép 1 Root/Thủy tổ (depth=0) duy nhất"),
    (22, "BUG-022", "TC_PER_030", "Sửa Person", "🟡 Medium", "Đổi gender Nam -> Nữ khi đang có vợ", "Popup Warning ảnh hưởng hôn nhân", "Cập nhật thẳng không cảnh báo", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "FE hiển thị cảnh báo xác nhận khi thay đổi giới tính của người đang có quan hệ SPOUSE"),
    (23, "BUG-023", "TC_REL_055", "Hôn nhân", "🟡 Medium", "Chèn quan hệ A kết hôn với chính A", "Chặn quan hệ tự thân (from_id == to_id)", "Thiếu check constraint", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "app/services/genealogy/genealogy_service.py: Chặn triệt để if from_person.id == to_person.id (HTTP 400)"),
    (24, "BUG-024", "TC_FAM_010", "Gán Trưởng chi", "🟡 Medium", "POST /branches kèm leaderPersonId", "Lưu Trưởng chi vào bảng branches", "HTTP 422 Extra field forbidden", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "app/controllers/member_api_router.py: BranchCreate schema hỗ trợ leaderPersonId và parentBranchId"),
    (25, "BUG-025", "TC_PER_036", "Xuất file Excel", "🟡 Medium", "GET /persons/export?format=csv", "Tải về file danh sách thành viên", "HTTP 404 Not Found", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "app/controllers/member_api_router.py: Thêm GET /families/{id}/persons/export xuất CSV chuẩn UTF-8 BOM"),
    (26, "BUG-026", "TC_FAM_005", "Upload Logo", "🟡 Medium", "Upload file logo_8MB.png", "HTTP 400 File vượt quá 5MB", "HTTP 404 Not Found", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "Client và Server chặn upload file kích thước > 5MB trước khi lưu trữ"),
    (27, "BUG-027", "TC_REL_053", "Tái hôn", "🟡 Medium", "Đã ly hôn, tạo lại kết hôn với người cũ", "Cho phép tạo bản ghi hôn nhân lần 2", "HTTP 400 Relationship already exists", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "Hệ thống hỗ trợ quan hệ hôn nhân nhiều giai đoạn dựa trên status='DIVORCED' và started_on"),
    (28, "BUG-028", "TC_SEC_088", "Rate Limiting", "🟡 Medium", "Bắn 100 requests API / 1 giây", "Từ request 61 trả về HTTP 429", "Xử lý toàn bộ 100 requests", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "app/core/rate_limit.py: Cấu hình In-Memory Sliding Window Rate Limiter chặn DDoS/Spam"),
    (29, "BUG-029", "TC_SEC_092", "Password Policy", "🟡 Medium", "Đổi mật khẩu thành '123456'", "HTTP 400 Yêu cầu >= 8 ký tự, hoa, số", "Chấp nhận mật khẩu yếu", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "app/schemas/common.py: Thêm validator kiểm tra >= 8 ký tự, có chữ hoa, chữ thường và chữ số"),
    (30, "BUG-030", "TC_PER_028", "Audit Log", "⚪ Low", "Sửa hồ sơ hoặc xóa quan hệ", "Chèn 1 bản ghi vào audit_logs", "Bảng audit_logs có 0 bản ghi", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "Ghi nhận vết thay đổi dữ liệu gia phả vào bảng audit_logs phục vụ truy vết"),
    (31, "BUG-031", "TC_SEC_091", "Bảo mật JSON API", "⚪ Low", "GET /api/v1/auth/me", "JSON không chứa từ khóa password", "Chứa requires_password_change", "✅ ĐÃ KHẮC PHỤC (RESOLVED)", "Kiểm soát an toàn thông tin nhạy cảm, không để lộ hash mật khẩu trong payload API"),
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Danh Sách Lỗi Sprint 2"

# Ensure grid lines visible
ws.views.sheetView[0].showGridLines = True

headers = [
    "STT", "Mã Lỗi", "Mã TC", "Chức Năng", "Mức Độ",
    "Dữ Liệu Đầu Vào (Input)", "Kết Quả Kỳ Vọng (Expected)",
    "Kết Quả Trước Khi Sửa (Actual Before)", "Trạng Thái (Status)",
    "Giải Pháp & Chi Tiết Sửa Lỗi"
]

header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

ws.append(headers)
for col_idx in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[1].height = 28

for row_idx, bug in enumerate(bugs_data, start=2):
    ws.append(list(bug))
    ws.row_dimensions[row_idx].height = 24
    
    # Fill zebra pattern
    is_even = (row_idx % 2 == 0)
    row_fill = PatternFill(start_color="F9FAFB" if is_even else "FFFFFF", fill_type="solid")
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.fill = row_fill
        cell.border = border_thin
        cell.font = Font(name="Calibri", size=10)
        
        # Center STT, Mã Lỗi, Mã TC, Mức độ, Trạng thái
        if col_idx in [1, 2, 3, 5, 9]:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            
        # Styling for Status column
        if col_idx == 9:
            cell.font = Font(name="Calibri", size=10, bold=True, color="008000")
            cell.fill = PatternFill(start_color="E6F4EA", fill_type="solid")
            
        # Severity colors
        if col_idx == 5:
            if "Critical" in str(bug[4]):
                cell.font = Font(name="Calibri", size=10, bold=True, color="9C0006")
            elif "High" in str(bug[4]):
                cell.font = Font(name="Calibri", size=10, bold=True, color="9C6500")

# Set column widths
col_widths = {
    1: 6,   # STT
    2: 12,  # Mã Lỗi
    3: 13,  # Mã TC
    4: 20,  # Chức Năng
    5: 14,  # Mức Độ
    6: 35,  # Input
    7: 35,  # Expected
    8: 35,  # Actual Before
    9: 25,  # Status
    10: 45  # Code fix
}

for col_idx, width in col_widths.items():
    col_letter = get_column_letter(col_idx)
    ws.column_dimensions[col_letter].width = width

wb.save(OUTPUT_FILE)
print(f"Successfully generated updated Excel report at: {OUTPUT_FILE}")
