import openpyxl
import json
import sys
import os

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
SPRINT2_DIR = os.path.dirname(CURRENT_DIR)

def get_file_path(filename, subfolder="test_cases"):
    candidates = [
        os.path.join(SPRINT2_DIR, subfolder, filename),
        os.path.join(CURRENT_DIR, filename),
        filename
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
    return os.path.join(SPRINT2_DIR, subfolder, filename)

def parse_and_classify():
    excel_path = get_file_path('MFGMS_AI_Sprint2_100_TestCases_Pro.xlsx')
    wb = openpyxl.load_workbook(excel_path)
    ws = wb['TestCases']
    
    headers = [cell.value for cell in ws[1]]
    
    test_cases = []
    
    for r in range(2, ws.max_row + 1):
        vals = [cell.value for cell in ws[r]]
        if not any(vals):
            continue
        
        tc = {
            'index': r - 1,
            'category': vals[0] or '',
            'code': vals[1] or '',
            'feature': vals[2] or '',
            'type_in_excel': vals[3] or '',
            'scenario': vals[4] or '',
            'input_action': vals[5] or '',
            'expected': vals[6] or '',
        }
        
        # Classification logic into:
        # 1. Security Test
        # 2. UI/UX Test (Frontend Browser)
        # 3. API Test (Backend)
        
        t_excel_lower = tc['type_in_excel'].lower()
        scenario_lower = tc['scenario'].lower()
        input_lower = tc['input_action'].lower()
        expected_lower = tc['expected'].lower()
        
        # Security criteria
        is_security = any([
            'security' in t_excel_lower,
            'xss' in t_excel_lower or 'xss' in scenario_lower or 'xss' in input_lower,
            'sql' in t_excel_lower or 'sqli' in t_excel_lower or 'sql injection' in scenario_lower or 'sql' in input_lower,
            'csrf' in t_excel_lower or 'csrf' in scenario_lower,
            'bảo mật' in t_excel_lower or 'bảo mật' in scenario_lower,
            'auth' in t_excel_lower and ('token giả' in input_lower or 'quyền' in scenario_lower or 'unauthorized' in expected_lower or '401' in expected_lower or '403' in expected_lower),
            'mã độc' in scenario_lower or 'mã độc' in input_lower,
            'phân quyền' in scenario_lower and ('chặn' in expected_lower or '403' in expected_lower),
            'idor' in t_excel_lower or 'idor' in scenario_lower,
            'tamper' in scenario_lower or 'sửa token' in input_lower,
            'traversal' in scenario_lower or 'lỗ hổng' in scenario_lower
        ])
        
        # UI/UX criteria
        is_ui = any([
            'ui' in t_excel_lower,
            'ux' in t_excel_lower,
            'giao diện' in t_excel_lower or 'giao diện' in scenario_lower or 'giao diện' in expected_lower,
            'responsive' in t_excel_lower or 'responsive' in scenario_lower,
            'modal' in scenario_lower or 'modal' in input_lower or 'popup' in scenario_lower or 'toast' in expected_lower,
            'hiển thị' in expected_lower and 'màu' in expected_lower,
            'form bị chặn ở fe' in expected_lower,
            'tooltip' in scenario_lower or 'hover' in scenario_lower or 'drag' in scenario_lower or 'kéo thả' in scenario_lower,
            'animation' in scenario_lower or 'zoom' in scenario_lower,
            'dark mode' in scenario_lower or 'theme' in scenario_lower,
            'nút' in scenario_lower or 'button' in scenario_lower or 'click' in input_lower or 'nhấn' in input_lower and 'giao diện' in expected_lower,
            'loading' in scenario_lower or 'skeleton' in scenario_lower or 'spinner' in expected_lower
        ])
        
        if is_security:
            classification = "Security Test"
        elif is_ui or ('fe' in expected_lower and 'api' not in expected_lower):
            classification = "UI/UX Test (Frontend Browser)"
        else:
            classification = "API Test (Backend)"
            
        tc['classification'] = classification
        test_cases.append(tc)
        
    out_path = get_file_path('classified_test_cases.json', 'test_cases')
    with open(out_path, 'w', encoding='utf-8') as f:
        json.dump(test_cases, f, ensure_ascii=False, indent=2)
        
    print(f"Total Test Cases: {len(test_cases)}")
    print(f"Saved to: {out_path}")
    
    counts = {}
    for tc in test_cases:
        c = tc['classification']
        counts[c] = counts.get(c, 0) + 1
        
    print("Classification counts:", counts)
    
if __name__ == '__main__':
    parse_and_classify()
